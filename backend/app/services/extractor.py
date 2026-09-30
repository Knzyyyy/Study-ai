import pymupdf  # Diubah dari import fitz
from pptx import Presentation
import os

def extract_text_from_pdf(file_path: str) -> list[dict]:
    """
    Mengekstrak teks dari PDF per halaman menggunakan PyMuPDF.
    Mengembalikan list of dictionary: [{"page_number": 1, "content_text": "..."}]
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File tidak ditemukan: {file_path}")

    results = []
    try:
        # Buka dokumen PDF
        doc = pymupdf.open(file_path)
        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            text = page.get_text("text").strip()
            
            # Jika teks kosong (mungkin hasil scan/gambar), berikan penanda
            if not text:
                text = "[Halaman ini kosong atau berisi gambar/hasil scan tanpa teks]"
                
            results.append({
                "page_number": page_num + 1,
                "content_text": text
            })
        doc.close()
    except Exception as e:
        raise Exception(f"Gagal memproses PDF: {str(e)}")
        
    return results

def extract_text_from_pptx(file_path: str) -> list[dict]:
    """
    Mengekstrak teks dari PPTX per slide menggunakan python-pptx.
    Termasuk teks dari Text Box, Tabel, dan Slide Notes.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File tidak ditemukan: {file_path}")

    results = []
    try:
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            slide_text_parts = []
            
            # 1. Ekstrak teks dari bentuk (shapes) dan tabel
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text_parts.append(shape.text.strip())
                
                # Cek apakah shape adalah tabel
                if shape.has_table:
                    for row in shape.table.rows:
                        row_data = []
                        for cell in row.cells:
                            if cell.text.strip():
                                row_data.append(cell.text.strip())
                        if row_data:
                            # Gabungkan cell dengan pemisah " | " agar formatnya cukup terbaca AI
                            slide_text_parts.append(" | ".join(row_data))
            
            # 2. Ekstrak teks dari Notes (Catatan Presenter)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_text_parts.append("\n[Notes]: " + notes)
            
            # Gabungkan semua teks di slide ini
            full_text = "\n".join(slide_text_parts).strip()
            
            if not full_text:
                full_text = "[Slide ini kosong atau hanya berisi gambar]"

            results.append({
                "page_number": i + 1,
                "content_text": full_text
            })
    except Exception as e:
        raise Exception(f"Gagal memproses PPTX: {str(e)}")
        
    return results